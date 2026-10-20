"""Extractor for agriculture/soybeans/Soybean_SD.xlsx (US WASDE + world FAS balance, house view)."""
from __future__ import annotations

import re

import openpyxl
import pandas as pd

from ..common import find_row, read_table


def _world_blocks(ws) -> pd.DataFrame:
    frames = []
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 1).value).strip() == "Country":
            title = str(ws.cell(r - 1, 1).value).strip()
            t = read_table(ws, "Country", start=r)
            t = t.loc[:, [c for c in t.columns if not c.startswith("RoW imports")]]
            t.insert(0, "block", title)
            frames.append(t)
    return pd.concat(frames, ignore_index=True)


def _grids(ws) -> dict[str, pd.DataFrame]:
    out = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and re.match(r"^\d\)", v):
            hdr = r + 1
            xs = [ws.cell(hdr, c).value for c in range(2, ws.max_column + 1)]
            xs = [x for x in xs if isinstance(x, (int, float))]
            rows = {}
            rr = hdr + 1
            while isinstance(ws.cell(rr, 1).value, (int, float)):
                rows[ws.cell(rr, 1).value] = [ws.cell(rr, 2 + i).value for i in range(len(xs))]
                rr += 1
            out[v.strip()] = pd.DataFrame.from_dict(rows, orient="index", columns=xs)
    return out


def _label_values(ws, first: str, last: str, col: int = 2) -> dict:
    a, b = find_row(ws, first), find_row(ws, last)
    return {str(ws.cell(r, 1).value).strip(): ws.cell(r, col).value
            for r in range(a, b + 1) if ws.cell(r, 1).value is not None}


def _text_block(ws, heading: str) -> list[str]:
    r = find_row(ws, heading) + 1
    lines = []
    while ws.cell(r, 1).value:
        lines.append(str(ws.cell(r, 1).value))
        r += 1
    return lines


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)

    us = read_table(wb["US_BS"], "Item")
    us = us[us["Unit"].notna()].reset_index(drop=True)
    value_cols = [c for c in us.columns if c not in ("Item", "Unit", "Source", "House - USDA Oct")]
    long = us.melt(id_vars=["Item", "Unit"], value_vars=value_cols, var_name="period", value_name="value")
    long = long[pd.to_numeric(long["value"], errors="coerce").notna()]
    long = long.rename(columns={"Item": "item", "Unit": "unit"})
    long["item"] = "US | " + long["item"].astype(str)

    world = _world_blocks(wb["World_BS"])
    wlong = world.melt(id_vars=["block", "Country"], var_name="field", value_name="value")
    wlong = wlong[pd.to_numeric(wlong["value"], errors="coerce").notna()]
    wlong = pd.DataFrame({"item": "World | " + wlong["Country"] + " | " + wlong["field"],
                          "unit": "MMT", "period": wlong["block"], "value": wlong["value"]})
    long = pd.concat([long, wlong], ignore_index=True)
    long["value"] = long["value"].astype(float)

    inputs = read_table(wb["Inputs"], "Parameter")
    inputs = inputs[inputs["Value"].notna()].reset_index(drop=True)

    pm = wb["Price_Model"]
    price = read_table(pm, "Case", label_col=2)

    grids = _grids(wb["Scenarios"])
    tables = {
        "long": long[["item", "unit", "period", "value"]],
        "us_bs": us,
        "world_bs": world,
        "key_metrics": read_table(wb["World_BS"], "Metric"),
        "inputs": inputs,
        "price_model": price,
        "price_hist": read_table(pm, "Marketing yr"),
        "spreads": read_table(wb["Spreads"], "Spread"),
    }
    for i, (title, g) in enumerate(grids.items(), 1):
        g = g.copy()
        g.index.name = "yield_bpa"
        tables[f"grid{i}"] = g.reset_index()

    crush_ws = wb["Crush"]
    meta = {
        "as_of": wb["README"]["A2"].value,
        "house_view": _text_block(wb["README"], "HOUSE VIEW"),
        "crush": _label_values(crush_ws, "Soybean meal ($/short ton)", "Oil share of product value"),
        "grid_titles": list(grids),
        "regression": {str(pm.cell(r, 2).value): pm.cell(r, 3).value
                       for r in range(find_row(pm, "REGRESSION"), find_row(pm, "IMPLIED 2026/27 PRICE"))},
    }
    return tables, meta
