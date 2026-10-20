"""Extractor for energy/oil/Oil_SD.xlsx (EIA STEO-based quarterly balance + scenarios)."""
from __future__ import annotations

import re

import openpyxl
import pandas as pd

from ..common import find_row, header_columns, read_table

PERIOD = re.compile(r"^(\d{2}Q[1-4]|\d{4})$")


def _balance(ws) -> pd.DataFrame:
    hr = find_row(ws, "Line item")
    cols = header_columns(ws, hr)
    periods = {h: c for h, c in cols.items() if PERIOD.match(h)}
    rows, section = [], ""
    for r in range(hr + 1, ws.max_row + 1):
        label, unit = ws.cell(r, 1).value, ws.cell(r, 2).value
        if label is None:
            continue
        if unit is None:                      # section heading (SUPPLY, DEMAND, ...)
            if str(label).isupper():
                section = str(label).strip()
            continue
        for p, c in periods.items():
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)):
                rows.append({"section": section, "item": f"{section} | {str(label).strip()}",
                             "label": str(label).strip(), "unit": unit, "period": p, "value": float(v)})
    return pd.DataFrame(rows)


def _indicators(ws) -> dict:
    r0 = find_row(ws, "INDICATORS")
    out = {}
    for r in range(r0 + 1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if label is None:
            continue
        v = ws.cell(r, 3).value
        if v is not None:
            out[str(label).strip()] = v
    return out


def _first_forecast(ws) -> str | None:
    """Forecast quarters are cream-filled in EIA_Data; return the first one."""
    hr = find_row(ws, "Series")
    r = find_row(ws, "World production")
    for h, c in header_columns(ws, hr).items():
        if PERIOD.match(h) and "Q" in h:
            rgb = ws.cell(r, c).fill.fgColor.rgb
            if isinstance(rgb, str) and rgb not in ("00000000", "FFFFFFFF"):
                return h
    return None


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    bal = _balance(wb["Balance"])

    sc_ws = wb["Scenarios"]
    start = find_row(sc_ws, "Scenario comparison")
    scen = read_table(sc_ws, "Scenario", start=start)
    scen = scen[scen["Prob."].notna()].reset_index(drop=True)

    tables = {
        "long": bal[["item", "unit", "period", "value"]],
        "balance": bal,
        "scenarios": scen,
        "agency_check": read_table(wb["Agency_Check"], "Metric"),
        "opec_plus": read_table(wb["OPEC+"], "Country"),
        "us_weekly": read_table(wb["US_Weekly"], "Week ending"),
        "catalysts": read_table(wb["Catalysts"], "When"),
    }
    readme = wb["README"]
    meta = {
        "as_of": readme["A2"].value,
        "active_scenario": sc_ws.cell(find_row(sc_ws, "Active scenario"), 3).value,
        "indicators": _indicators(wb["Balance"]),
        "forecast_from": _first_forecast(wb["EIA_Data"]),
    }
    return tables, meta
