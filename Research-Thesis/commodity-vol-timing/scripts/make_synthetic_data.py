"""Create a synthetic dataset with the same layout as the Bloomberg export.

The numbers are simulated (stochastic volatility with occasional jumps, a
time-varying futures curve whose generics roll on the exchange calendar, and
a small implied-volatility premium). They are NOT market data: use them to
check that the code runs end to end without a Bloomberg licence.

    python scripts/make_synthetic_data.py --commodity GC   # -> data/raw/synthetic_GC.xlsx
    python src/run_all.py --commodity GC --data data/raw/synthetic_GC.xlsx
"""
import argparse
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from commodities import SPECS, bloomberg_ticker, contract_calendar  # noqa: E402


def simulate(code="CL", start="2007-01-02", end="2026-09-30", seed=0, price0=60.0):
    rng = np.random.default_rng(seed)
    dates = pd.bdate_range(start, end)
    n = len(dates)

    # log-variance: AR(1) with rare upward jumps (crisis regimes)
    log_v = np.empty(n)
    log_v[0] = np.log(0.30 ** 2)
    for t in range(1, n):
        jump = rng.random() < 0.002
        log_v[t] = (log_v[t - 1] + 0.02 * (np.log(0.30 ** 2) - log_v[t - 1])
                    + 0.12 * rng.standard_normal() + (1.2 if jump else 0.0))
    vol = np.sqrt(np.exp(log_v))
    spot = price0 * np.exp(np.cumsum(vol / np.sqrt(252) * rng.standard_normal(n)))

    # futures curve: F = S * exp(carry * time to expiry); carry drifts between
    # contango (>0) and backwardation (<0). Generics roll after each last trade.
    carry = 0.05 + 0.10 * np.sin(np.arange(n) / 400) + 0.01 * rng.standard_normal(n)
    ltd = pd.DatetimeIndex(contract_calendar(code, dates)["last_trade"])
    pos = np.arange(n)
    ltd_pos = np.searchsorted(dates.values, ltd.values)
    k = np.searchsorted(ltd_pos, pos, side="left")
    k1 = np.minimum(k, len(ltd_pos) - 1)
    k2 = np.minimum(k + 1, len(ltd_pos) - 1)
    tau1 = np.maximum(ltd_pos[k1] - pos, 0) / 252
    tau2 = np.maximum(ltd_pos[k2] - pos, 21) / 252
    basis = 0.04 * rng.standard_normal(len(ltd_pos) + 1)   # contract-specific offset
    f1 = spot * np.exp(carry * tau1 + basis[k1])
    f2 = spot * np.exp(carry * tau2 + basis[k2])

    iv30 = 100 * vol * (1.04 + 0.05 * rng.standard_normal(n))      # small premium
    iv1m = iv30 * (1 + 0.03 * rng.standard_normal(n))
    iv2m = iv30 * (0.98 + 0.03 * rng.standard_normal(n))
    return pd.DataFrame({
        "date": dates, "f1": f1, "open": f1 * np.exp(rng.normal(0, 0.005, n)),
        "high": f1 * np.exp(np.abs(rng.normal(0, 0.01, n))),
        "low": f1 * np.exp(-np.abs(rng.normal(0, 0.01, n))),
        "volume": rng.integers(50_000, 400_000, n), "f2": f2,
        "iv30": iv30, "iv1m": iv1m, "iv2m": iv2m})


def write_bloomberg_layout(df, path, code="CL"):
    """Columns A-F (F1), I-J (F2), L-O (implied vol), as in a BDH export."""
    wb = Workbook()
    ws = wb.active
    t1, t2 = bloomberg_ticker(code, 1), bloomberg_ticker(code, 2)
    ws["A1"], ws["I1"], ws["L1"] = t1, t2, t1
    for col, name in zip("BCDEF", ["PX_SETTLE", "PX_OPEN", "PX_HIGH", "PX_LOW", "PX_VOLUME"]):
        ws[f"{col}2"] = name
    ws["J2"] = "PX_SETTLE"
    for col, name in zip("MNO", ["30DAY_IMPVOL_100.0%MNY_DF", "1ST_MTH_IMPVOL_100.0%MNY_DF",
                                 "2ND_MTH_IMPVOL_100.0%MNY_DF"]):
        ws[f"{col}2"] = name
    for i, r in enumerate(df.itertuples(index=False), start=3):
        d = r.date.to_pydatetime()
        ws.cell(i, 1, d)
        for j, v in enumerate([r.f1, r.open, r.high, r.low], start=2):
            ws.cell(i, j, round(float(v), 4))
        ws.cell(i, 6, int(r.volume))
        ws.cell(i, 9, d)
        ws.cell(i, 10, round(float(r.f2), 4))
        ws.cell(i, 12, d)
        for j, v in enumerate([r.iv30, r.iv1m, r.iv2m], start=13):
            ws.cell(i, j, round(float(v), 2))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Write a synthetic dataset in Bloomberg layout.")
    ap.add_argument("--commodity", default="CL", choices=list(SPECS))
    ap.add_argument("--out", help="default: data/raw/synthetic_CODE.xlsx")
    ap.add_argument("--seed", type=int, default=0)
    a = ap.parse_args()
    out = Path(a.out) if a.out else ROOT / "data" / "raw" / f"synthetic_{a.commodity}.xlsx"
    write_bloomberg_layout(simulate(a.commodity, seed=a.seed), out, a.commodity)
    print(f"Synthetic {SPECS[a.commodity]['name']} data written to {out}")
